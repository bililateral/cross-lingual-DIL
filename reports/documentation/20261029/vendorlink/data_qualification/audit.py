"""One-shot authorized inventory, Linux py310 only; no views/qrels/models created."""
import collections
import hashlib
import json
import os
from pathlib import Path
import re
import resource
import time
import unicodedata

from openpyxl import load_workbook

ROOT = Path('/home/yongpeng/cross-lingual')
OUT = ROOT / 'reports/documentation/20261029/vendorlink/data_qualification'
resource.setrlimit(resource.RLIMIT_AS, (4*1024**3, 4*1024**3))
resource.setrlimit(resource.RLIMIT_CPU, (600, 600))
started = time.monotonic()


def canonical(s):
    return ' '.join(unicodedata.normalize('NFKC', s).split())


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024*1024), b''):
            h.update(block)
    return h.hexdigest()


targets = {'alphabay': 'AlphaBay', 'alpha bay': 'AlphaBay',
           'dream': 'Dream', 'dream market': 'Dream',
           'silk road': 'Silk Road', 'silkroad': 'Silk Road'}
report = {'scope': 'source_schema_and_target_account_eligibility_only',
          'observed_server': time.strftime('%Y-%m-%dT%H:%M:%S%z'),
          'no_model_load': True, 'no_views_or_qrels_written': True,
          'target_name_mapping': targets, 'files': {}, 'archive_directories': {},
          'excluded_contents': ['vendor/PGP SQL and HTML identity sources',
                                'private_custody', 'existing evaluation truth/results'],
          'source_code_sha256': digest(Path(__file__))}
for name in ('market_item.xlsx', '2017-12-05-philipjames11-darknetmarketplacedataagora20142015.xlsx',
             'products_data.csv', 'tijkc3xx.sql', '3z669jwe.sql'):
    p = ROOT/name
    report['files'][name] = {'exists': p.exists(), 'bytes': p.stat().st_size if p.exists() else None}
for name in ('html-rips', '2017-12-09-grams-christopher'):
    paths = [p for p in (ROOT/name).iterdir() if p.is_file()]
    report['archive_directories'][name] = {'files': len(paths), 'bytes': sum(p.stat().st_size for p in paths),
                                         'contents_read': False, 'role': 'vendor_identity_snapshots_not_ad_table'}
source = ROOT/'market_item.xlsx'
before = source.stat()
report['files'][source.name]['sha256'] = digest(source)
book = load_workbook(source, read_only=True, data_only=True)
sheet = book.worksheets[0]
headers = list(next(sheet.iter_rows(max_row=1, values_only=True)))
assert headers == ['vendor', 'ship_from', 'title', 'description', 'price', 'category', 'market'], headers
report['fields'] = headers
report['field_mapping'] = {k: k for k in ('market', 'vendor', 'title', 'description')}
market_counts = collections.Counter()
stats = {m: {'raw_rows': 0, 'invalid_rows': 0, 'invalid_fields': collections.Counter(),
             'valid_rows': 0, 'raw_accounts': set(), 'valid_accounts': set()}
         for m in ('AlphaBay', 'Dream', 'Silk Road')}
accounts, raw_owners, owners, raw_fingerprints = {}, {}, {}, {}
rows = 0
for row in sheet.iter_rows(min_row=2, values_only=True):
    rows += 1
    vendor, _, title, description, _, _, market = row
    market_name = canonical(market) if isinstance(market, str) else '<invalid_market>'
    market_counts[market_name] += 1
    target = targets.get(market_name.casefold())
    if target is None:
        continue
    s = stats[target]
    s['raw_rows'] += 1
    if isinstance(vendor, str) and canonical(vendor):
        s['raw_accounts'].add(canonical(vendor))
    values = {'vendor': vendor, 'title': title, 'description': description}
    invalid = [k for k, v in values.items() if not isinstance(v, str) or not canonical(v)]
    if invalid:
        s['invalid_rows'] += 1
        s['invalid_fields'].update(invalid)
        continue
    vendor, title, description = (canonical(v) for v in (vendor, title, description))
    key = (target, vendor)
    s['valid_rows'] += 1
    s['valid_accounts'].add(vendor)
    raw = hashlib.sha256((title+'\0'+description).casefold().encode()).hexdigest()
    raw_owners.setdefault(raw, set()).add(key)
    pattern = re.compile(r'(?<!\w)'+re.escape(vendor)+r'(?!\w)', re.IGNORECASE)
    masked = hashlib.sha256((pattern.sub('[VENDOR]', title)+'\0'+pattern.sub('[VENDOR]', description)).casefold().encode()).hexdigest()
    owners.setdefault(masked, set()).add(key)
    accounts.setdefault(key, set()).add(masked)
    raw_fingerprints.setdefault((key, masked), set()).add(raw)
book.close()
after = source.stat()
assert (before.st_size, before.st_mtime_ns) == (after.st_size, after.st_mtime_ns), 'Source changed during audit'
for market, s in stats.items():
    lengths, unique, shared = [], 0, 0
    for key, fs in accounts.items():
        if key[0] != market:
            continue
        clean = [f for f in fs if len(owners[f]) == 1 and all(len(raw_owners[r]) == 1 for r in raw_fingerprints[(key, f)])]
        lengths.append(len(clean))
        unique += len(fs)
        shared += len(fs)-len(clean)
    s.update(raw_accounts=len(s['raw_accounts']), valid_accounts=len(s['valid_accounts']),
             within_account_unique_masked_ads=unique, cross_account_shared_removed=shared,
             clean_ads=sum(lengths), eligible_accounts_ge20=sum(n >= 20 for n in lengths),
             clean_ads_in_eligible_accounts=sum(n for n in lengths if n >= 20))
report.update(rows=rows, rows_by_market=dict(market_counts.most_common()), target_markets=stats,
              qualification='TARGET_THREE_MARKET_AD_TABLES_NOT_PRESENT' if not sum(s['raw_rows'] for s in stats.values()) else 'INSPECT_TARGET_COUNTS',
              wall_seconds=time.monotonic()-started, cpu_seconds=time.process_time(),
              peak_rss_kib=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
(OUT/'inventory.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'status': report['qualification'], 'rows': rows, 'target_markets': stats,
                  'wall_seconds': report['wall_seconds']}, ensure_ascii=False))
